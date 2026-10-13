import os, sqlite3, sys
from pathlib import Path
os.chdir(Path(__file__).resolve().parent)  # caminhos relativos à raiz do projeto
import numpy as np, pandas as pd
sys.path.insert(0, "src")
import atraso as A
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

# ============================================================ 1. DADOS DE BASE (mesmo escopo dos notebooks)
D = "powerbi/dados/"
ped = pd.read_csv(D + "fato_pedidos.csv", parse_dates=["data_compra"])
itens = pd.read_csv(D + "fato_itens.csv"); prod = pd.read_csv(D + "dim_produtos.csv")
vend = pd.read_csv(D + "dim_vendedores.csv"); cli = pd.read_csv(D + "dim_clientes.csv"); uf = pd.read_csv(D + "dim_uf.csv")
ped["ym"] = ped.data_compra.dt.strftime("%Y-%m")

mensal = ped.groupby("ym").agg(pedidos=("order_id", "size"), receita=("receita_itens", "sum"),
          entregues=("flag_atrasado", "count"), atrasados=("flag_atrasado", "sum"),
          avaliacoes=("review_score", "count"), soma_notas=("review_score", "sum")).reset_index()

u = (ped.merge(uf[["uf", "regiao"]], on="uf").groupby(["uf", "regiao"])
       .agg(pedidos=("order_id", "size"), receita=("receita_itens", "sum"), frete=("frete_total", "sum"),
            entregues=("flag_atrasado", "count"), atrasados=("flag_atrasado", "sum"), soma_dias=("dias_entrega", "sum"),
            avaliacoes=("review_score", "count"), soma_notas=("review_score", "sum")).reset_index())
u["_ord"] = u.atrasados / u.entregues; u = u.sort_values("_ord", ascending=False).reset_index(drop=True)

sat = ped.dropna(subset=["review_score", "dias_atraso"]).copy()
sat["faixa"] = pd.cut(sat.dias_atraso, [-np.inf, 0, 3, 7, np.inf], labels=["No prazo", "2. Atraso 1-3 dias", "3. Atraso 4-7 dias", "4. Atraso > 7 dias"])
faixas = sat.groupby("faixa", observed=True).agg(pedidos=("order_id", "size"), soma=("review_score", "sum"),
                                                  baixas=("review_score", lambda s: int((s <= 2).sum()))).reset_index()

cat = (itens.merge(prod, on="product_id").groupby("categoria").price.sum().sort_values(ascending=False).head(10).reset_index())
seg = cli.groupby("segmento").agg(clientes=("customer_unique_id", "size"), receita=("monetario", "sum")).sort_values("receita", ascending=False).reset_index()
rv = itens.groupby("seller_id").price.sum().sort_values(ascending=False).values
q = np.array_split(rv, 5)
quint = pd.DataFrame({"quintil": [1, 2, 3, 4, 5], "vendedores": [len(x) for x in q], "receita": [x.sum() for x in q]})

# modelo (mesmas funções do notebook 02)
con = sqlite3.connect("olist.db"); dd = A.montar_dataset(con)
res, prev = A.avaliar_janelas(dd, sorted(dd.ym.unique())[8:])
bins = [-np.inf, 0, 5, 10, 15, np.inf]; rot = ["Folga de 0 dias ou menos (prazo menor que o típico)", "Folga de 0 a 5 dias", "Folga de 5 a 10 dias", "Folga de 10 a 15 dias", "Folga acima de 15 dias"]
prev["faixa"] = pd.cut(prev.folga_r, bins, labels=rot)
fol = prev.groupby("faixa", observed=True).agg(pedidos=("y", "size"), atrasados=("y", "sum")).reset_index()
cap = []
for k in (0.05, 0.10, 0.20, 0.30):
    pm = prev.groupby("ym").apply(lambda x: pd.Series({"c": A.captura_topk(x.y, x.p, k),
         "p": x.y.values[np.argsort(-x.p.values)[:int(len(x) * k)]].mean(), "b": x.y.mean()}))
    cap.append((k, pm.c.mean(), pm.p.mean(), (pm.p / pm.b).mean()))
modelos = (res.groupby("modelo").agg(auc=("auc", "mean"), dp=("auc", "std"), pior=("auc", "min"), cap10=("cap10", "mean"), lift10=("lift10", "mean"))
           .sort_values("auc", ascending=False).reset_index())

# ============================================================ 2. ESTILOS
F = "Arial"
f_base = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF")                     # entrada (valor calculado fora da planilha)
f_link = Font(name=F, size=10, color="008000")                   # vínculo com outra aba
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=16, bold=True, color="1E293B"); f_sub = Font(name=F, size=10, italic=True, color="64748B")
f_sec = Font(name=F, size=11, bold=True, color="1E293B")
fill_head = PatternFill("solid", fgColor="2563EB"); fill_tot = PatternFill("solid", fgColor="E2E8F0"); fill_key = PatternFill("solid", fgColor="EFF6FF")
thin = Side(style="thin", color="CBD5E1"); border = Border(bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top"); WRAPI = Alignment(wrap_text=True, vertical="top", indent=1); CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
BRL = '"R$" #,##0;("R$" #,##0);"-"'; BRL2 = '"R$" #,##0.00'; INT = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; DEC1 = '0.0'; DEC2 = '0.00'; MULT = '0.0"x"'

wb = Workbook(); wb.remove(wb.active)
def sheet(nome, titulo, sub, larguras):
    ws = wb.create_sheet(nome); ws.sheet_view.showGridLines = False
    ws["A1"] = titulo; ws["A1"].font = f_title; ws["A2"] = sub; ws["A2"].font = f_sub
    for i, w in enumerate(larguras, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def head(ws, row, cols):
    for j, t in enumerate(cols, 1):
        c = ws.cell(row, j, t); c.font = f_head; c.fill = fill_head; c.alignment = CTR
    ws.row_dimensions[row].height = 32
def put(ws, ref, v, fmt=None, font=None, fill=None, al=None):
    c = ws[ref]; c.value = v; c.font = font or f_base
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    return c
def totrow(ws, row, ncol):
    for j in range(1, ncol + 1):
        c = ws.cell(row, j); c.fill = fill_tot; c.font = Font(name=F, size=10, bold=True, color=c.font.color.rgb if c.font and c.font.color and c.font.color.rgb not in (None, "00000000") and c.font.color.rgb == "00008000" else "000000")

# ============================================================ 3. ABAS DE DADOS (antes do Resumo, que as referencia)
# ---- KPIs_Mensais
ws = sheet("KPIs_Mensais", "KPIs mensais", "Pedidos entregues, compra entre jan/2017 e ago/2018. Azul = dado de entrada; preto = fórmula.", [11, 7, 7, 11, 15, 13, 14, 12, 10, 12, 13, 11, 12])
head(ws, 4, ["Ano-Mês", "Ano", "Mês", "Pedidos", "Receita (R$)", "Receita por Pedido (R$)", "Entregues c/ data", "Atrasados", "% Atraso", "Avaliações", "Soma das notas", "Nota Média", "Receita MoM %"])
r0 = 5
for i, m in mensal.iterrows():
    r = r0 + i
    put(ws, f"A{r}", m.ym, font=f_in); put(ws, f"B{r}", int(m.ym[:4]), font=f_in); put(ws, f"C{r}", int(m.ym[5:]), font=f_in)
    put(ws, f"D{r}", int(m.pedidos), INT, f_in); put(ws, f"E{r}", round(float(m.receita), 2), BRL, f_in)
    put(ws, f"F{r}", f"=E{r}/D{r}", BRL2); put(ws, f"G{r}", int(m.entregues), INT, f_in); put(ws, f"H{r}", int(m.atrasados), INT, f_in)
    put(ws, f"I{r}", f"=H{r}/G{r}", PCT); put(ws, f"J{r}", int(m.avaliacoes), INT, f_in); put(ws, f"K{r}", int(m.soma_notas), INT, f_in)
    put(ws, f"L{r}", f"=K{r}/J{r}", DEC2)
    if i > 0: put(ws, f"M{r}", f"=E{r}/E{r-1}-1", PCT)
rl = r0 + len(mensal) - 1; rt = rl + 1
put(ws, f"A{rt}", "Total")
for col in "DEGHJK": put(ws, f"{col}{rt}", f"=SUM({col}{r0}:{col}{rl})", BRL if col == "E" else INT)
put(ws, f"F{rt}", f"=E{rt}/D{rt}", BRL2); put(ws, f"I{rt}", f"=H{rt}/G{rt}", PCT); put(ws, f"L{rt}", f"=K{rt}/J{rt}", DEC2)
for j in range(1, 14): ws.cell(rt, j).fill = fill_tot; ws.cell(rt, j).font = f_bold
ws.freeze_panes = "B5"
KM = dict(r0=r0, rl=rl, rt=rt)

# ---- Logistica_UF
ws = sheet("Logistica_UF", "Logística por estado e região", "Entrega x prazo estimado. UFs com menos de 200 pedidos têm amostra pequena. Azul = dado de entrada; preto = fórmula.",
           [7, 14, 10, 10, 13, 12, 10, 16, 3, 15, 13, 12, 11, 13, 11, 12])
head(ws, 4, ["UF", "Região", "Pedidos", "% Atraso", "Entrega Média (dias)", "Frete % Receita", "Nota Média", "Amostra", "", "Receita (R$)", "Frete (R$)", "Entregues c/ data", "Atrasados", "Soma dias entrega", "Avaliações", "Soma das notas"])
ws["I4"].fill = PatternFill(fill_type=None)
r0 = 5
for i, x in u.iterrows():
    r = r0 + i
    put(ws, f"A{r}", x.uf, font=f_in); put(ws, f"B{r}", x.regiao, font=f_in); put(ws, f"C{r}", int(x.pedidos), INT, f_in)
    put(ws, f"D{r}", f"=M{r}/L{r}", PCT); put(ws, f"E{r}", f"=N{r}/L{r}", DEC1); put(ws, f"F{r}", f"=K{r}/J{r}", PCT)
    put(ws, f"G{r}", f"=P{r}/O{r}", DEC2); put(ws, f"H{r}", f'=IF(C{r}<200,"amostra pequena","ok")')
    put(ws, f"J{r}", round(float(x.receita), 2), BRL, f_in); put(ws, f"K{r}", round(float(x.frete), 2), BRL, f_in)
    put(ws, f"L{r}", int(x.entregues), INT, f_in); put(ws, f"M{r}", int(x.atrasados), INT, f_in); put(ws, f"N{r}", round(float(x.soma_dias), 1), INT, f_in)
    put(ws, f"O{r}", int(x.avaliacoes), INT, f_in); put(ws, f"P{r}", int(x.soma_notas), INT, f_in)
rl = r0 + len(u) - 1; rt = rl + 1
put(ws, f"A{rt}", "Total"); put(ws, f"C{rt}", f"=SUM(C{r0}:C{rl})", INT)
for col in "JKLMNOP": put(ws, f"{col}{rt}", f"=SUM({col}{r0}:{col}{rl})", BRL if col in "JK" else INT)
put(ws, f"D{rt}", f"=M{rt}/L{rt}", PCT); put(ws, f"E{rt}", f"=N{rt}/L{rt}", DEC1); put(ws, f"F{rt}", f"=K{rt}/J{rt}", PCT); put(ws, f"G{rt}", f"=P{rt}/O{rt}", DEC2)
for j in list(range(1, 9)) + list(range(10, 17)): ws.cell(rt, j).fill = fill_tot; ws.cell(rt, j).font = f_bold
LU = dict(r0=r0, rl=rl, rt=rt)
rg0 = rt + 3
put(ws, f"A{rg0-1}", "Por região", font=f_sec); head(ws, rg0, ["Região", "", "Pedidos", "% Atraso", "Entrega Média (dias)", "Frete % Receita", "Nota Média"])
ws[f"B{rg0}"].fill = PatternFill(fill_type=None)
regs = ["Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"]; REG = {}
for i, rg in enumerate(regs):
    r = rg0 + 1 + i; REG[rg] = r; rng = lambda c: f"{c}${r0}:{c}${rl}"
    put(ws, f"A{r}", rg); put(ws, f"C{r}", f"=SUMIFS({rng('C')},{rng('B')},$A{r})", INT)
    put(ws, f"D{r}", f"=SUMIFS({rng('M')},{rng('B')},$A{r})/SUMIFS({rng('L')},{rng('B')},$A{r})", PCT)
    put(ws, f"E{r}", f"=SUMIFS({rng('N')},{rng('B')},$A{r})/SUMIFS({rng('L')},{rng('B')},$A{r})", DEC1)
    put(ws, f"F{r}", f"=SUMIFS({rng('K')},{rng('B')},$A{r})/SUMIFS({rng('J')},{rng('B')},$A{r})", PCT)
    put(ws, f"G{r}", f"=SUMIFS({rng('P')},{rng('B')},$A{r})/SUMIFS({rng('O')},{rng('B')},$A{r})", DEC2)
ws.freeze_panes = "C5"
ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "AL e MA concentram os maiores atrasos"; ch.legend = None
ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=rl), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=rl))
ch.y_axis.number_format = "0%"; ch.x_axis.scaling.orientation = "maxMin"; ch.height = 13; ch.width = 14
ch.x_axis.delete = False; ch.y_axis.delete = False
ch.series[0].graphicalProperties.solidFill = "2563EB"; ws.add_chart(ch, f"J{rg0-1}")

# ---- Atraso_vs_Nota
ws = sheet("Atraso_vs_Nota", "Atraso x satisfação do cliente", "Pedidos entregues com data de entrega e avaliação. Azul = dado de entrada; preto = fórmula.", [30, 12, 15, 17, 12, 14, 14, 20])
head(ws, 4, ["Faixa de atraso", "Pedidos", "Soma das notas", "Notas baixas (1 a 2)", "Nota Média", "% Notas Baixas", "% dos pedidos", "% das notas baixas totais"])
for i, x in faixas.iterrows():
    r = 5 + i
    put(ws, f"A{r}", str(x.faixa), font=f_in); put(ws, f"B{r}", int(x.pedidos), INT, f_in); put(ws, f"C{r}", int(x.soma), INT, f_in); put(ws, f"D{r}", int(x.baixas), INT, f_in)
    put(ws, f"E{r}", f"=C{r}/B{r}", DEC2); put(ws, f"F{r}", f"=D{r}/B{r}", PCT); put(ws, f"G{r}", f"=B{r}/B$9", PCT); put(ws, f"H{r}", f"=D{r}/D$9", PCT)
put(ws, "A9", "Total"); put(ws, "B9", "=SUM(B5:B8)", INT); put(ws, "C9", "=SUM(C5:C8)", INT); put(ws, "D9", "=SUM(D5:D8)", INT)
put(ws, "E9", "=C9/B9", DEC2); put(ws, "F9", "=D9/B9", PCT); put(ws, "G9", "=B9/B$9", PCT); put(ws, "H9", "=D9/D$9", PCT)
for j in range(1, 9): ws.cell(9, j).fill = fill_tot; ws.cell(9, j).font = f_bold
put(ws, "A12", "Indicadores derivados", font=f_sec)
put(ws, "A13", "% dos pedidos que chegaram com atraso"); put(ws, "B13", "=SUM(B6:B8)/B9", PCT)
put(ws, "A14", "% das notas baixas vindas de pedidos atrasados"); put(ws, "B14", "=SUM(D6:D8)/D9", PCT)
put(ws, "A15", "Queda de nota: no prazo vs atraso > 7 dias"); put(ws, "B15", "=E5-E8", DEC2)
ch = BarChart(); ch.type = "col"; ch.style = 10; ch.title = "Satisfação por condição de entrega"; ch.legend = None
ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=8), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=8))
ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 5; ch.height = 7.5; ch.width = 14; ch.x_axis.delete = False; ch.y_axis.delete = False
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False; ch.dataLabels.showLegendKey = False; ch.dataLabels.showPercent = False
ch.series[0].graphicalProperties.solidFill = "2563EB"; ws.add_chart(ch, "A18")

# ---- Categorias
ws = sheet("Categorias", "Receita por categoria: Top 10", "Receita de itens (pedidos entregues, jan/2017 a ago/2018). Azul = dado de entrada; verde = vínculo com outra aba; preto = fórmula.", [32, 16, 14, 14])
head(ws, 4, ["Categoria", "Receita (R$)", "% da receita", "% acumulado"])
for i, x in cat.iterrows():
    r = 5 + i; put(ws, f"A{r}", x.categoria, font=f_in); put(ws, f"B{r}", round(float(x.price), 2), BRL, f_in)
    put(ws, f"C{r}", f"=B{r}/B$16", PCT); put(ws, f"D{r}", f"=SUM(C$5:C{r})", PCT)
put(ws, "A15", "Demais categorias"); put(ws, "B15", "=B16-SUM(B5:B14)", BRL); put(ws, "C15", "=B15/B$16", PCT); put(ws, "D15", "=SUM(C$5:C15)", PCT)
put(ws, "A16", "Total (= receita da aba KPIs_Mensais)"); put(ws, "B16", f"=KPIs_Mensais!E{KM['rt']}", BRL, f_link); put(ws, "C16", "=B16/B$16", PCT)
for j in range(1, 5): ws.cell(16, j).fill = fill_tot; ws.cell(16, j).font = Font(name=F, size=10, bold=True, color="008000" if j == 2 else "000000")
put(ws, "A18", "Categorias no top 10 respondem por"); put(ws, "B18", "=D14", PCT)

# ---- Clientes_RFM
ws = sheet("Clientes_RFM", "Clientes: recompra e segmentação RFM", "Segmentos por regras de negócio (recência, valor, recorrência) sobre o período inteiro. Azul = dado de entrada; preto = fórmula.", [30, 12, 16, 13, 12, 14])
head(ws, 4, ["Segmento", "Clientes", "Receita (R$)", "% Clientes", "% Receita", "Receita por cliente (R$)"])
n = len(seg)
for i, x in seg.iterrows():
    r = 5 + i; put(ws, f"A{r}", x.segmento, font=f_in); put(ws, f"B{r}", int(x.clientes), INT, f_in); put(ws, f"C{r}", round(float(x.receita), 2), BRL, f_in)
    put(ws, f"D{r}", f"=B{r}/B${5+n}", PCT); put(ws, f"E{r}", f"=C{r}/C${5+n}", PCT); put(ws, f"F{r}", f"=C{r}/B{r}", BRL2)
rt = 5 + n; put(ws, f"A{rt}", "Total"); put(ws, f"B{rt}", f"=SUM(B5:B{rt-1})", INT); put(ws, f"C{rt}", f"=SUM(C5:C{rt-1})", BRL)
put(ws, f"D{rt}", f"=B{rt}/B{rt}", PCT); put(ws, f"E{rt}", f"=C{rt}/C{rt}", PCT); put(ws, f"F{rt}", f"=C{rt}/B{rt}", BRL2)
for j in range(1, 7): ws.cell(rt, j).fill = fill_tot; ws.cell(rt, j).font = f_bold
CR = dict(rt=rt, r1=5, r2=rt - 1)
put(ws, f"A{rt+3}", "Indicadores derivados", font=f_sec)
put(ws, f"A{rt+4}", "Taxa de recompra (clientes com 2+ pedidos)"); put(ws, f"B{rt+4}", f'=INDEX(B5:B{rt-1},MATCH("Recorrentes",A5:A{rt-1},0))/B{rt}', PCT)
put(ws, f"A{rt+5}", "% da receita: Inativos de alto valor"); put(ws, f"B{rt+5}", f'=INDEX(E5:E{rt-1},MATCH("Inativos de alto valor",A5:A{rt-1},0))', PCT)
put(ws, f"A{rt+6}", "Clientes: Inativos de alto valor"); put(ws, f"B{rt+6}", f'=INDEX(B5:B{rt-1},MATCH("Inativos de alto valor",A5:A{rt-1},0))', INT)
put(ws, f"A{rt+7}", "Receita por cliente: Inativos de alto valor"); put(ws, f"B{rt+7}", f'=INDEX(F5:F{rt-1},MATCH("Inativos de alto valor",A5:A{rt-1},0))', BRL2)
CRX = dict(recompra=f"B{rt+4}", inat_rec=f"B{rt+5}", inat_cli=f"B{rt+6}", inat_tkt=f"B{rt+7}")

# ---- Vendedores
ws = sheet("Vendedores", "Concentração de receita por vendedor", "Vendedores ordenados por receita e divididos em 5 grupos de igual tamanho (quintis). Azul = dado de entrada; preto = fórmula.", [26, 13, 16, 14, 12, 14])
head(ws, 4, ["Quintil (1 = maiores)", "Vendedores", "Receita (R$)", "% vendedores", "% Receita", "% receita acumulado"])
for i, x in quint.iterrows():
    r = 5 + i; put(ws, f"A{r}", int(x.quintil), font=f_in); put(ws, f"B{r}", int(x.vendedores), INT, f_in); put(ws, f"C{r}", round(float(x.receita), 2), BRL, f_in)
    put(ws, f"D{r}", f"=B{r}/B$10", PCT); put(ws, f"E{r}", f"=C{r}/C$10", PCT); put(ws, f"F{r}", f"=SUM(E$5:E{r})", PCT)
put(ws, "A10", "Total"); put(ws, "B10", "=SUM(B5:B9)", INT); put(ws, "C10", "=SUM(C5:C9)", BRL); put(ws, "D10", "=B10/B10", PCT); put(ws, "E10", "=C10/C10", PCT)
for j in range(1, 7): ws.cell(10, j).fill = fill_tot; ws.cell(10, j).font = f_bold
put(ws, "A13", "Os 20% maiores vendedores respondem por"); put(ws, "B13", "=E5", PCT)
put(ws, "A14", "Os 40% menores vendedores respondem por"); put(ws, "B14", "=E8+E9", PCT)

# ---- Modelo_Risco
ws = sheet("Modelo_Risco", "Modelo de risco de atraso (resultados fora do tempo)", "Fonte: notebook 02. 12 janelas móveis (set/2017 a ago/2018), 74.206 pedidos. Azul = dado de entrada do notebook; preto = fórmula.", [58, 14, 14, 14, 14, 14])
put(ws, "A3", "1. Taxa de atraso por folga do prazo (prazo prometido − tempo típico da rota)", font=f_sec)
head(ws, 4, ["Folga do prazo", "Pedidos", "Atrasados", "Taxa de atraso", "% dos pedidos"])
for i, x in fol.iterrows():
    r = 5 + i; put(ws, f"A{r}", str(x.faixa), font=f_in); put(ws, f"B{r}", int(x.pedidos), INT, f_in); put(ws, f"C{r}", int(x.atrasados), INT, f_in)
    put(ws, f"D{r}", f"=C{r}/B{r}", PCT); put(ws, f"E{r}", f"=B{r}/B$10", PCT)
put(ws, "A10", "Total"); put(ws, "B10", "=SUM(B5:B9)", INT); put(ws, "C10", "=SUM(C5:C9)", INT); put(ws, "D10", "=C10/B10", PCT); put(ws, "E10", "=B10/B10", PCT)
for j in range(1, 6): ws.cell(10, j).fill = fill_tot; ws.cell(10, j).font = f_bold
put(ws, "A11", "Risco relativo: folga <= 0 dias vs folga > 15 dias"); put(ws, "B11", "=D5/D9", MULT)
put(ws, "A13", "2. Captura de atrasos ao sinalizar os X% de maior risco (média dos 12 meses)", font=f_sec)
head(ws, 14, ["% de pedidos sinalizados", "Atrasos capturados", "Precisão (% dos sinalizados que atrasam)", "Lift vs. taxa-base"])
for i, (k, c, p, l) in enumerate(cap):
    r = 15 + i; put(ws, f"A{r}", k, "0%", f_in); put(ws, f"B{r}", round(float(c), 4), PCT, f_in); put(ws, f"C{r}", round(float(p), 4), PCT, f_in); put(ws, f"D{r}", round(float(l), 3), MULT, f_in)
put(ws, "A20", "3. Comparação de modelos (média dos 12 meses de teste; AUC 0,5 = acaso)", font=f_sec)
head(ws, 21, ["Modelo", "AUC médio", "Desvio-padrão do AUC", "AUC do pior mês", "Captura no top 10%", "Lift no top 10%"])
for i, x in modelos.iterrows():
    r = 22 + i; put(ws, f"A{r}", x.modelo, font=f_in); put(ws, f"B{r}", round(float(x.auc), 3), DEC2, f_in); put(ws, f"C{r}", round(float(x.dp), 3), DEC2, f_in)
    put(ws, f"D{r}", round(float(x.pior), 3), DEC2, f_in); put(ws, f"E{r}", round(float(x.cap10), 4), PCT, f_in); put(ws, f"F{r}", round(float(x.lift10), 3), MULT, f_in)
for j in range(1, 7): ws.cell(22, j).fill = fill_key
put(ws, "A29", "Linha destacada = modelo escolhido (regressão logística com folga da rota, janela de 90 dias).", font=f_sub)

# ============================================================ 4. RESUMO, RECOMENDAÇÕES, NOTAS
KMt = f"KPIs_Mensais"; rtm = KM["rt"]
ws = sheet("Resumo", "Olist: resumo executivo", "Marketplace brasileiro | pedidos entregues, compra de jan/2017 a ago/2018 | dados reais públicos (Kaggle). Todos os números abaixo são fórmulas ligadas às demais abas.", [46, 18, 70])
put(ws, "A4", "Indicadores-chave", font=f_sec); head(ws, 5, ["Indicador", "Valor", "Leitura"])
ya = f"INDEX(KPIs_Mensais!M{KM['r0']}:M{KM['rl']},MATCH(\"2017-11\",KPIs_Mensais!A{KM['r0']}:A{KM['rl']},0))"
kpis = [
 ("Receita (itens, R$)", f"=KPIs_Mensais!E{rtm}", BRL, "Receita de produtos, sem frete e sem margem (a base não tem custos)."),
 ("Pedidos entregues", f"=KPIs_Mensais!D{rtm}", INT, "Escopo: status entregue, 20 meses completos."),
 ("Receita por Pedido (R$)", f"=KPIs_Mensais!F{rtm}", BRL2, "Estável ao longo do período: o crescimento vem de volume."),
 ("Receita YoY % (jan-ago/2018 vs jan-ago/2017)", f'=SUMIFS(KPIs_Mensais!E{KM["r0"]}:E{KM["rl"]},KPIs_Mensais!B{KM["r0"]}:B{KM["rl"]},2018,KPIs_Mensais!C{KM["r0"]}:C{KM["rl"]},"<=8")/SUMIFS(KPIs_Mensais!E{KM["r0"]}:E{KM["rl"]},KPIs_Mensais!B{KM["r0"]}:B{KM["rl"]},2017,KPIs_Mensais!C{KM["r0"]}:C{KM["rl"]},"<=8")-1', PCT, "Mesmos oito meses de cada ano, para comparação justa."),
 ("% Atraso", f"=KPIs_Mensais!I{rtm}", PCT, "Atraso medido contra a data estimada informada ao cliente."),
 ("Nota Média (1 a 5)", f"=KPIs_Mensais!L{rtm}", DEC2, "Média geral; veja a queda por faixa de atraso abaixo."),
 ("Taxa de Recompra %", f"=Clientes_RFM!{CRX['recompra']}", PCT, "Clientes com 2 ou mais pedidos: retenção é o maior espaço de crescimento."),
 ("Receita Top 20% Vendedores %", "=Vendedores!B13", PCT, "Receita concentrada em poucos vendedores."),
]
for i, (nm, fm, nf, lt) in enumerate(kpis):
    r = 6 + i; put(ws, f"A{r}", nm, font=f_bold); c = put(ws, f"B{r}", fm, nf, f_link); c.alignment = Alignment(horizontal="right"); c.font = Font(name=F, size=11, bold=True, color="008000")
    put(ws, f"C{r}", lt, al=WRAPI)
    for j in "ABC": ws[f"{j}{r}"].border = border
r = 6 + len(kpis) + 1
put(ws, f"A{r}", "Principais achados", font=f_sec); head(ws, r + 1, ["Achado", "Valor", "Indicador"])
ach = [
 ("Atraso derruba a nota", f"=Atraso_vs_Nota!E5", DEC2, "Nota média dos pedidos entregues no prazo"),
 ("", f"=Atraso_vs_Nota!E8", DEC2, "Nota média dos pedidos com atraso acima de 7 dias"),
 ("Poucos pedidos atrasados concentram a insatisfação", "=Atraso_vs_Nota!B13", PCT, "% dos pedidos com atraso, entre os que têm avaliação (difere levemente do indicador geral)"),
 ("", "=Atraso_vs_Nota!B14", PCT, "% das notas baixas (1 a 2) vindas desses pedidos"),
 ("Nordeste sofre mais com a entrega", f"=Logistica_UF!D{REG['Nordeste']}", PCT, "% de atraso no Nordeste"),
 ("", f"=Logistica_UF!D{REG['Sudeste']}", PCT, "% de atraso no Sudeste (referência)"),
 ("Prazo apertado para a rota atrasa muito mais", "=Modelo_Risco!D5", PCT, "Taxa de atraso quando o prazo prometido é igual ou menor que o típico da rota"),
 ("", "=Modelo_Risco!D9", PCT, "Taxa de atraso quando há folga de mais de 15 dias"),
 ("Dá para antecipar parte dos atrasos na compra", "=Modelo_Risco!B16", PCT, "% dos atrasos capturados ao sinalizar os 10% de pedidos de maior risco"),
 ("Black Friday exige capacidade", f"={ya}", PCT, "Variação da receita de nov/2017 sobre out/2017"),
]
for i, (a, fm, nf, ind) in enumerate(ach):
    rr = r + 2 + i; put(ws, f"A{rr}", a, font=f_bold, al=WRAP); c = put(ws, f"B{rr}", fm, nf); c.font = Font(name=F, size=11, bold=True, color="008000"); c.alignment = Alignment(horizontal="right", vertical="top")
    put(ws, f"C{rr}", ind, al=WRAPI)
    for j in "ABC": ws[f"{j}{rr}"].border = border
rch = r + 2 + len(ach) + 1
put(ws, f"A{rch}", "Gráficos", font=f_sec)
c1 = BarChart(); c1.type = "col"; c1.style = 10; c1.title = "Receita por Ano-Mês (R$)"; c1.legend = None
c1.add_data(Reference(wb["KPIs_Mensais"], min_col=5, min_row=4, max_row=KM["rl"]), titles_from_data=True)
c1.set_categories(Reference(wb["KPIs_Mensais"], min_col=1, min_row=KM["r0"], max_row=KM["rl"])); c1.y_axis.number_format = '#,##0'
c1.height = 7.5; c1.width = 12; c1.x_axis.delete = False; c1.y_axis.delete = False
c1.series[0].graphicalProperties.solidFill = "2563EB"; ws.add_chart(c1, f"A{rch+1}")
c2 = BarChart(); c2.type = "col"; c2.style = 10; c2.title = "Satisfação por condição de entrega"; c2.legend = None
c2.add_data(Reference(wb["Atraso_vs_Nota"], min_col=5, min_row=4, max_row=8), titles_from_data=True); c2.set_categories(Reference(wb["Atraso_vs_Nota"], min_col=1, min_row=5, max_row=8))
c2.y_axis.scaling.min = 0; c2.y_axis.scaling.max = 5; c2.height = 7.5; c2.width = 13; c2.x_axis.delete = False; c2.y_axis.delete = False
c2.dataLabels = DataLabelList(); c2.dataLabels.showVal = True; c2.dataLabels.showSerName = False; c2.dataLabels.showCatName = False; c2.dataLabels.showLegendKey = False; c2.dataLabels.showPercent = False
c2.series[0].graphicalProperties.solidFill = "2563EB"; ws.add_chart(c2, f"C{rch+1}")
from openpyxl.worksheet.pagebreak import Break
ws.row_breaks.append(Break(id=rch - 1))

# ---- Recomendacoes
ws = sheet("Recomendacoes", "Recomendações priorizadas", "Prioridade = julgamento do analista a partir da evidência. Valores são fórmulas ligadas às abas de dados (verde).", [4, 11, 16, 46, 40, 40, 13, 13])
head(ws, 4, ["#", "Prioridade", "Área", "Ação recomendada", "Evidência", "Métrica-chave (valor atual | referência)", "Valor atual", "Referência"])
recs = [
 ("Alta", "Operações / Atendimento", "Tratar atrasos acima de 3 dias como incidente: aviso proativo ao cliente, atendimento prioritário e compensação.", "A nota despenca nos atrasos mais longos, onde a maioria das avaliações é 1 ou 2.", "Nota média com atraso > 7 dias | nota média no prazo", "=Atraso_vs_Nota!E8", DEC2, "=Atraso_vs_Nota!E5", DEC2),
 ("Alta", "Logística", "Atacar a logística do Nordeste (AL, MA, PI, CE, SE, BA): revisar transportadoras e rotas; avaliar centros de distribuição.", "Maiores taxas de atraso, prazos mais longos e frete mais pesado.", "% de atraso no Nordeste | Sudeste", f"=Logistica_UF!D{REG['Nordeste']}", PCT, f"=Logistica_UF!D{REG['Sudeste']}", PCT),
 ("Alta", "Logística / Produto", "Calibrar o prazo prometido pela rota: onde o prazo é igual ou menor que o típico, estender a promessa ou reforçar a operação. Testar em A/B o efeito na conversão.", "Prazo apertado para a rota atrasa muito mais; o prazo bruto sozinho não é bom preditor.", "Atraso com folga <= 0 | folga > 15 dias", "=Modelo_Risco!D5", PCT, "=Modelo_Risco!D9", PCT),
 ("Alta", "Operações / Dados", "Triagem proativa: sinalizar o top 10% de risco para aviso ao cliente, priorização no despacho e acompanhamento.", "Os pedidos de maior risco concentram boa parte dos atrasos.", "Atrasos capturados no top 10% | sorteio aleatório", "=Modelo_Risco!B16", PCT, "=Modelo_Risco!A16", PCT),
 ("Média", "Comercial / CRM", "Reativar clientes inativos de alto valor e fazer onboarding dos novos de alto valor.", "A recompra é muito baixa e não é explicada por atraso ou nota baixa; o gargalo é relacionamento.", "% da receita nos inativos de alto valor | taxa de recompra", f"=Clientes_RFM!{CRX['inat_rec']}", PCT, f"=Clientes_RFM!{CRX['recompra']}", PCT),
 ("Média", "Comercial", "Reduzir a dependência de poucos vendedores: programa de desenvolvimento do 2º quintil e monitoramento de risco nos maiores.", "A receita está concentrada nos maiores vendedores.", "% da receita dos 20% maiores | 20% proporcional", "=Vendedores!B13", PCT, "=Vendedores!D5", PCT),
 ("Média", "Operações", "Planejar capacidade para picos sazonais (Black Friday).", "Pico expressivo de receita e de atraso em nov/2017.", "Receita nov/17 sobre out/17 | média mensal de variação", f"={ya}", PCT, f"=AVERAGE(KPIs_Mensais!M{KM['r0']+1}:M{KM['rl']})", PCT),
 ("Média", "Dados", "Retreinar o modelo todo mês e monitorar o AUC; investigar se cair perto de 0,6.", "A política de prazos mudou em 2018 e relações do passado deixaram de valer (concept drift).", "AUC médio | AUC do pior mês", "=Modelo_Risco!B22", DEC2, "=Modelo_Risco!D22", DEC2),
]
for i, (pr, ar, ac, ev, mt, v1, f1, v2, f2) in enumerate(recs):
    r = 5 + i
    put(ws, f"A{r}", i + 1, al=Alignment(vertical="top", horizontal="center")); put(ws, f"B{r}", pr, font=Font(name=F, size=10, bold=True, color="DC2626" if pr == "Alta" else "B45309"), al=WRAP)
    put(ws, f"C{r}", ar, al=WRAP); put(ws, f"D{r}", ac, al=WRAP); put(ws, f"E{r}", ev, al=WRAP); put(ws, f"F{r}", mt, al=WRAP)
    c = put(ws, f"G{r}", v1, f1, f_link); c.alignment = Alignment(horizontal="right", vertical="top"); c = put(ws, f"H{r}", v2, f2, f_link); c.alignment = Alignment(horizontal="right", vertical="top")
    ws.row_dimensions[r].height = 66
    for j in "ABCDEFGH": ws[f"{j}{r}"].border = border
ws.freeze_panes = "A5"

# ---- Notas
ws = sheet("Notas", "Notas, premissas e limitações", "Leia antes de usar os números.", [28, 110])
linhas = [
 ("Fonte", "Brazilian E-Commerce Public Dataset by Olist (Kaggle, olistbr/brazilian-ecommerce). Dados reais, ~100 mil pedidos."),
 ("Escopo", "Pedidos com status entregue e compra entre jan/2017 e ago/2018 (meses completos). Cliente = customer_unique_id. Receita = soma do preço dos itens, sem frete."),
 ("Como esta planilha foi feita", "Os dados de base (células em AZUL) foram calculados em SQL/Python a partir da base original e colados como valores. Todo o resto (PRETO) é fórmula; vínculos entre abas aparecem em VERDE."),
 ("Reprodução", "Scripts do projeto: carregar_banco.py, sql/01_views.sql, sql/02_consultas_negocio.sql, exportar_powerbi.py e notebooks 01 e 02. A planilha é regenerada por build_xlsx.py."),
 ("Atraso", "Pedido entregue depois da data estimada informada ao cliente. 8 pedidos entregues sem data de entrega ficam fora das métricas de prazo."),
 ("Nota Média", "Média das avaliações dos pedidos. Pedidos sem avaliação ficam fora do denominador."),
 ("RFM", "Segmentação por regras de negócio (recência, valor e recorrência), não por clusters estatísticos. Como ~97% dos clientes compraram uma única vez, quintis de frequência não funcionam."),
 ("Modelo de risco", "Usa só informações conhecidas na compra. Resultados vêm de validação temporal em 12 janelas móveis. A versão final (janela de 90 dias para o tempo típico da rota) foi escolhida depois de ver os resultados, o que gera leve otimismo."),
 ("Limitação 1", "Análise observacional: associação não é causalidade (atraso e nota baixa andam juntos, mas outros problemas também pesam na nota)."),
 ("Limitação 2", "Receita não é lucro: a base não tem custo nem margem."),
 ("Limitação 3", "Janela de 20 meses de uma empresa em crescimento acelerado; tendências podem refletir a expansão e não comportamento estável."),
 ("Limitação 4", "O modelo tem desempenho moderado e probabilidades não calibradas: use-o para ordenar risco (triagem), não como probabilidade absoluta."),
 ("Prioridades", "A coluna Prioridade da aba Recomendacoes é um julgamento do analista, não resultado de cálculo."),
]
head(ws, 4, ["Tema", "Nota"])
for i, (a, b) in enumerate(linhas):
    r = 5 + i; put(ws, f"A{r}", a, font=f_bold, al=WRAP); put(ws, f"B{r}", b, al=WRAP); ws.row_dimensions[r].height = 30
    for j in "AB": ws[f"{j}{r}"].border = border

# ordem das abas
ordem = ["Resumo", "Recomendacoes", "KPIs_Mensais", "Logistica_UF", "Atraso_vs_Nota", "Categorias", "Clientes_RFM", "Vendedores", "Modelo_Risco", "Notas"]
wb._sheets = [wb[n] for n in ordem]
wb["Resumo"].sheet_properties.tabColor = "2563EB"; wb["Recomendacoes"].sheet_properties.tabColor = "DC2626"
from openpyxl.worksheet.properties import PageSetupProperties
for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"; _ws.page_setup.fitToWidth = 1; _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    _ws.print_options.horizontalCentered = True
wb.active = 0
os.makedirs("excel", exist_ok=True)
wb.save("excel/resumo_executivo_olist.xlsx")
print("salvo")
